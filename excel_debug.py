from openpyxl import load_workbook

wb = load_workbook(
    "resources/templates/Спецификация.xlsx"
)

print("=" * 80)

print("Листы:")

for ws in wb.worksheets:
    print(" -", ws.title)

print()

ws = wb["Спецификация"]

print("=" * 80)
print("Объединённые диапазоны")
print("=" * 80)

for rng in ws.merged_cells.ranges:
    print(rng)

print()

print("=" * 80)
print("Размер листа")
print("=" * 80)

print("max_row =", ws.max_row)
print("max_column =", ws.max_column)

print()

print("=" * 80)
print("Непустые ячейки")
print("=" * 80)

for row in ws.iter_rows():

    for cell in row:

        if cell.value is not None:

            print(
                f"{cell.coordinate:6}",
                repr(cell.value)
            )